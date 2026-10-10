from openpyxl import Workbook, load_workbook
from datetime import datetime
import os

class ExcelManager:
    filename = "inventory.xlsx"
    headers = ["Product ID","Product Name","Category","Quantity","Price","Supplier","Reorder Level","Last Updated"]

    def __init__(self):
        if not os.path.exists(self.filename):
            wb = Workbook()
            ws = wb.active
            ws.title = "Inventory"
            ws.append(self.headers)
            ws.freeze_panes = "A2"
            widths = [16,25,18,12,12,25,16,22]
            for i,w in enumerate(widths,1):
                ws.column_dimensions[chr(64+i)].width = w
            wb.save(self.filename)

    def _read(self):
        wb = load_workbook(self.filename)
        ws = wb["Inventory"]
        return list(ws.iter_rows(min_row=2, values_only=True))

    def get_all(self):
        return [list(x) for x in self._read() if x[0] is not None]

    def exists(self, product_id):
        return any(str(x[0]) == str(product_id) for x in self.get_all())

    def add(self, p):
        wb = load_workbook(self.filename)
        ws = wb["Inventory"]
        ws.append([p["Product ID"],p["Product Name"],p.get("Category",""),p["Quantity"],
                   p["Price"],p.get("Supplier",""),p["Reorder Level"],
                   datetime.now().strftime("%Y-%m-%d %H:%M:%S")])
        wb.save(self.filename)

    def update(self, old_id, p):
        wb = load_workbook(self.filename)
        ws = wb["Inventory"]
        for row in range(2, ws.max_row+1):
            if str(ws.cell(row,1).value) == str(old_id):
                vals = [p["Product ID"],p["Product Name"],p.get("Category",""),p["Quantity"],
                        p["Price"],p.get("Supplier",""),p["Reorder Level"],
                        datetime.now().strftime("%Y-%m-%d %H:%M:%S")]
                for col,val in enumerate(vals,1):
                    ws.cell(row,col).value = val
                break
        wb.save(self.filename)

    def delete(self, product_id):
        wb = load_workbook(self.filename)
        ws = wb["Inventory"]
        for row in range(2, ws.max_row+1):
            if str(ws.cell(row,1).value) == str(product_id):
                ws.delete_rows(row,1)
                break
        wb.save(self.filename)
